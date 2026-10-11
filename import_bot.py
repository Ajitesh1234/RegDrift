"""Turn a real RPA configuration file into a bot entry in bots.json, so RegDrift checks real settings.

UiPath (REFramework):   python import_bot.py Data/Config.xlsx --name "KYC-Refresh-Bot" --description "..."
Any other RPA tool:     put the bot's settings in a CSV with columns Name,Value[,Description]
                        (e.g. Blue Prism environment variables) and run:
                        python import_bot.py settings.csv --name "STR-Filer-Bot" --platform "Blue Prism" --description "..."

Reads the Settings, Constants and Assets sheets of a UiPath Config.xlsx (columns: Name, Value, Description).
Names are converted to snake_case (KYCReverificationYears -> kyc_reverification_years).
Values: numbers, true/false, and lists written as  a;b;c  or  a|b|c  or a JSON array.
Re-running updates the same bot (matched by name), so it is safe to repeat.
"""
import argparse, csv, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
SHEETS = ("Settings", "Constants", "Assets")


def snake(name: str) -> str:
    s = re.sub(r"(?<=[a-z0-9])(?=[A-Z])|(?<=[A-Z])(?=[A-Z][a-z])", "_", str(name).strip())
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def parse_value(v):
    if v is None:
        return None
    if isinstance(v, bool) or isinstance(v, (int, float)):
        return int(v) if isinstance(v, float) and v.is_integer() else v
    s = str(v).strip()
    if not s:
        return None
    if s.lower() in ("true", "false"):
        return s.lower() == "true"
    if s.startswith("[") and s.endswith("]"):
        try:
            return json.loads(s)
        except ValueError:
            pass
    for sep in (";", "|"):
        if sep in s:
            return [p.strip() for p in s.split(sep) if p.strip()]
    try:
        f = float(s)
        return int(f) if f.is_integer() and "." not in s else f
    except ValueError:
        return s


def read_settings(path):
    """Return ({snake_name: value}, [warnings])."""
    rows, warnings = [], []
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = [s for s in SHEETS if s in wb.sheetnames]
        if not sheets:
            raise ValueError(f"none of the sheets {SHEETS} found in {path} (found: {wb.sheetnames})")
        for s in sheets:
            it = wb[s].iter_rows(values_only=True)
            header = [str(h).strip().lower() if h is not None else "" for h in next(it, [])]
            if "name" not in header or "value" not in header:
                warnings.append(f"sheet {s}: no Name/Value columns, skipped")
                continue
            ni, vi = header.index("name"), header.index("value")
            rows += [(s, r[ni], r[vi]) for r in it if r and len(r) > max(ni, vi)]
    elif path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as fh:
            for r in csv.DictReader(fh):
                low = {k.strip().lower(): v for k, v in r.items() if k}
                rows.append(("csv", low.get("name"), low.get("value")))
    else:
        raise ValueError("expected a .xlsx (UiPath Config.xlsx) or .csv file")
    params = {}
    for sheet, name, value in rows:
        if not name or not str(name).strip():
            continue
        v, key = parse_value(value), snake(name)
        if v is None or not key:
            continue
        if key in params and params[key] != v:
            warnings.append(f"duplicate setting {key!r} ({sheet}); keeping the later value")
        params[key] = v
    return params, warnings


def upsert(bots, entry):
    """Update the bot with the same name (case-insensitive), else append with a fresh id."""
    for i, b in enumerate(bots):
        if b["name"].strip().lower() == entry["name"].strip().lower():
            entry = {**b, **{k: v for k, v in entry.items() if v not in (None, "", [])}, "id": b["id"]}
            bots[i] = entry
            return bots, "updated", entry
    nums = [int(m.group(1)) for b in bots if (m := re.fullmatch(r"BOT-(\d+)", b["id"]))]
    entry = {"id": entry.get("id") or f"BOT-{max(nums, default=0) + 1:03d}", **{k: v for k, v in entry.items() if k != "id"}}
    bots.append(entry)
    return bots, "added", entry


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("config", help="UiPath Config.xlsx or a Name,Value CSV")
    ap.add_argument("--name", required=True, help="bot name, e.g. KYC-Refresh-Bot")
    ap.add_argument("--description", default="", help="what the bot does (used by the RAG search: write a real sentence)")
    ap.add_argument("--platform", default="UiPath")
    ap.add_argument("--tags", default="", help="comma-separated keywords")
    ap.add_argument("--bots", default=os.path.join(HERE, "bots.json"))
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args(argv)
    try:
        params, warnings = read_settings(a.config)
    except (OSError, ValueError) as e:
        print("ERROR:", e)
        return 1
    for w in warnings:
        print("warning:", w)
    if not params:
        print("ERROR: no settings found")
        return 1
    if not a.description.strip():
        print("warning: no --description; the RAG search matches on descriptions, so add one for better results")
    with open(a.bots, encoding="utf-8") as fh:
        bots = json.load(fh)
    entry = {"name": a.name, "platform": a.platform, "description": a.description.strip(),
             "tags": [t.strip() for t in a.tags.split(",") if t.strip()], "parameters": params}
    bots, what, entry = upsert(bots, entry)
    print(f"{what} {entry['id']} {entry['name']} with {len(params)} settings: {', '.join(list(params)[:8])}{' ...' if len(params) > 8 else ''}")
    if not a.dry_run:
        tmp = a.bots + ".tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(bots, fh, indent=2)
        os.replace(tmp, a.bots)
    return 0


if __name__ == "__main__":
    sys.exit(main())
