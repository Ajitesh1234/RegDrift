import csv, json, os, sys, tempfile, unittest
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
for k in ("GROQ_API_KEY", "LORA_ADAPTER", "JIRA_URL"): os.environ.pop(k, None)
try:
    from openpyxl import Workbook
except ImportError:  # optional dependency: skip these tests instead of crashing the run
    raise unittest.SkipTest("openpyxl not installed (pip install -r requirements.txt)")
import import_bot, regdrift


def make_config(path, extra_settings=()):
    wb = Workbook(); ws = wb.active; ws.title = "Settings"
    ws.append(["Name", "Value", "Description"])
    ws.append(["OrchestratorQueueName", "KYC_Queue", "queue"])
    ws.append(["KYCReverificationYearsHighRisk", 3, "years between re-verification"])
    ws.append(["SendReminderEmails", "TRUE", ""])
    for r in extra_settings: ws.append(list(r))
    c = wb.create_sheet("Constants"); c.append(["Name", "Value", "Description"])
    c.append(["AcceptedCustomerTypes", "NRI;PIO", "list"]); c.append(["MaxRetryNumber", 2.0, ""]); c.append(["EmptyOne", None, ""])
    wb.create_sheet("Other").append(["x"])
    wb.save(path)


class Snake(unittest.TestCase):
    def test_names(self):
        self.assertEqual(import_bot.snake("KYCReverificationYearsHighRisk"), "kyc_reverification_years_high_risk")
        self.assertEqual(import_bot.snake("MaxRetryNumber"), "max_retry_number")
        self.assertEqual(import_bot.snake("str_filing_days"), "str_filing_days")
        self.assertEqual(import_bot.snake(" Sanctions Screening (days) "), "sanctions_screening_days")

    def test_values(self):
        p = import_bot.parse_value
        self.assertEqual([p("3"), p(3.0), p("2.5"), p("TRUE"), p("a;b|c"), p('["x","y"]'), p("  "), p(None), p("text")],
                         [3, 3, 2.5, True, ["a", "b|c"] if False else ["a", "b|c"], ["x", "y"], None, None, "text"])


class Import(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()
        self.xlsx = os.path.join(self.d, "Config.xlsx"); make_config(self.xlsx)
        self.bots = os.path.join(self.d, "bots.json")
        with open(self.bots, "w") as fh: json.dump([{"id": "BOT-007", "name": "Other", "platform": "UiPath", "description": "d", "tags": [], "parameters": {}}], fh)

    def load(self):
        with open(self.bots) as fh: return json.load(fh)

    def test_reads_all_sheets_with_types(self):
        params, warn = import_bot.read_settings(self.xlsx)
        self.assertEqual(params["kyc_reverification_years_high_risk"], 3)
        self.assertIs(params["send_reminder_emails"], True)
        self.assertEqual(params["accepted_customer_types"], ["NRI", "PIO"])
        self.assertEqual(params["max_retry_number"], 2)
        self.assertNotIn("empty_one", params)

    def test_add_then_update_is_idempotent(self):
        args = [self.xlsx, "--name", "KYC-Refresh-Bot", "--description", "Refreshes KYC", "--tags", "kyc, refresh", "--bots", self.bots]
        self.assertEqual(import_bot.main(args), 0)
        first = self.load()
        self.assertEqual([b["id"] for b in first], ["BOT-007", "BOT-008"])
        self.assertEqual(import_bot.main(args), 0)
        self.assertEqual(self.load(), first)                          # same bot updated, nothing duplicated
        self.assertEqual(first[1]["tags"], ["kyc", "refresh"])

    def test_dry_run_writes_nothing(self):
        before = self.load()
        self.assertEqual(import_bot.main([self.xlsx, "--name", "X", "--bots", self.bots, "--dry-run"]), 0)
        self.assertEqual(self.load(), before)

    def test_csv_for_other_platforms(self):
        p = os.path.join(self.d, "bp.csv")
        with open(p, "w", newline="") as fh:
            w = csv.writer(fh); w.writerow(["Name", "Value"]); w.writerow(["STRFilingDays", "7"])
        self.assertEqual(import_bot.main([p, "--name", "STR-Filer-Bot", "--platform", "Blue Prism", "--bots", self.bots]), 0)
        b = self.load()[-1]
        self.assertEqual((b["platform"], b["parameters"]), ("Blue Prism", {"str_filing_days": 7}))

    def test_bad_inputs_fail_cleanly(self):
        bad = os.path.join(self.d, "x.docx"); open(bad, "w").close()
        self.assertEqual(import_bot.main([bad, "--name", "X", "--bots", self.bots]), 1)
        self.assertEqual(import_bot.main([os.path.join(self.d, "nope.xlsx"), "--name", "X", "--bots", self.bots]), 1)

    def test_imported_bot_is_checked_by_the_pipeline_end_to_end(self):
        import_bot.main([self.xlsx, "--name", "KYC-Refresh-Bot", "--description", "Refreshes customer KYC", "--bots", self.bots])
        os.environ["REGDRIFT_BOTS"] = self.bots
        orig = regdrift.llm_json
        regdrift.llm_json = lambda p: [{"clause": "(1)", "summary": "KYC re-verification every 2 years", "parameter": "kyc_reverification_years_high_risk",
                                        "operator": "<=", "required_value": 2, "unit": "years"}] if "Extract machine" in p else None
        try:
            r = regdrift.run({"circular_text": "RBI/2026-27/5\n(1) High-risk customers shall be re-verified at least every 2 years.\n", "dry_run": True})
        finally:
            regdrift.llm_json = orig; os.environ.pop("REGDRIFT_BOTS")
        t = r["tickets"][0]
        self.assertEqual((t["bot"], t["status"]), ("KYC-Refresh-Bot", "violating"))
        self.assertIn("3", t["reason"])


if __name__ == "__main__":
    unittest.main()
